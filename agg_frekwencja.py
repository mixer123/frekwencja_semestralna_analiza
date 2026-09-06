import sys, csv
import openpyxl
from collections import defaultdict
from datetime import datetime

rok = sys.argv[1]  # np. 2020_2021
infile = f"dane/frekw{rok}.xlsx"
outfile = f"wyniki_semestralne/agg_{rok}.csv"

wb = openpyxl.load_workbook(infile, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]

agg = defaultdict(int)  # (dziennik, data_iso, wpis_nazwa) -> count

for i, r in enumerate(ws.iter_rows(values_only=True)):
    if i == 0:
        continue
    dziennik = r[1]
    data = r[3]
    wpis = r[9]
    if data is None:
        continue
    key = (dziennik, data.strftime("%Y-%m-%d"), wpis)
    agg[key] += 1

with open(outfile, "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(["Dziennik", "Data", "Wpis", "Liczba"])
    for (dziennik, data, wpis), cnt in agg.items():
        w.writerow([dziennik, data, wpis, cnt])

print(f"OK {rok}: {len(agg)} agg rows z {i+1} wierszy zrodlowych -> {outfile}")
