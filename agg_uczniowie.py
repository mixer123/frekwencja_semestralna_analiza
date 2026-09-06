import sys, csv, hashlib
import openpyxl
from collections import defaultdict
import pandas as pd

rok = sys.argv[1]  # np. 2020_2021
infile = f"dane/frekw{rok}.xlsx"
outfile = f"wyniki_semestralne/agg_uczniowie_{rok}.csv"

SALT = "moja_szkola_2026"

def normalizuj(tekst):
    return " ".join(str(tekst).strip().split()).upper()

def parsuj_ucznia(uczen_str):
    return str(uczen_str).split(",")[0].strip()

def uczen_hash(nazwisko_imie, szkola):
    klucz = f"{SALT}|{normalizuj(szkola)}|{normalizuj(nazwisko_imie)}"
    return hashlib.sha256(klucz.encode()).hexdigest()[:10].upper()

kolumny_nieobecnosci = {'nieob. uspr. p.s.', 'nieob. uspraw.', 'nieobecność'}
kolumny_obecnosci = {'nauka zdalna', 'obecność', 'spóźn. uspr.', 'spóźnienie', 'zwolniony'}
wszystkie_wpisy = kolumny_nieobecnosci | kolumny_obecnosci

rok_szkolny = rok.replace('_', '/')
rok_start = rok_szkolny.split('/')[0]
rok_koniec = rok_szkolny.split('/')[1]
odciecie_wrzesien = pd.Timestamp(f"{rok_start}-09-01")
koniec_kwietnia = pd.Timestamp(f"{rok_koniec}-04-30")

GRANICE = {
    '2020/2021': ('2020-12-22', '2021-01-18'),
    '2021/2022': ('2022-01-28', '2022-02-14'),
    '2022/2023': ('2023-02-10', '2023-02-27'),
    '2023/2024': ('2024-01-12', '2024-01-29'),
    '2024/2025': ('2024-12-20', '2025-01-07'),
    '2025/2026': ('2025-12-19', '2026-01-07'),
}
koniec_sem1 = pd.Timestamp(GRANICE[rok_szkolny][0])
poczatek_sem2 = pd.Timestamp(GRANICE[rok_szkolny][1])

wb = openpyxl.load_workbook(infile, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]

agg = defaultdict(lambda: [0, 0])  # (uczen_hash, dziennik, semestr) -> [wszystkich, nieobecnych]
hash_cache = {}

import re
def numer_klasy(dziennik):
    m = re.match(r"^(\d+)", str(dziennik))
    return int(m.group(1)) if m else None

for i, r in enumerate(ws.iter_rows(values_only=True)):
    if i == 0:
        continue
    szkola = r[0]
    dziennik = r[1]
    uczen = r[2]
    data = r[3]
    wpis = r[9]
    if data is None or wpis not in wszystkie_wpisy:
        continue
    if data <= odciecie_wrzesien:
        continue
    if data <= koniec_sem1:
        semestr = 'I'
    elif data >= poczatek_sem2:
        semestr = 'II'
    else:
        continue  # ferie - poza semestrami
    if numer_klasy(dziennik) == 5 and semestr == 'II' and data > koniec_kwietnia:
        continue

    ck = (uczen, szkola)
    h = hash_cache.get(ck)
    if h is None:
        h = uczen_hash(parsuj_ucznia(uczen), szkola)
        hash_cache[ck] = h

    key = (h, dziennik, semestr)
    slot = agg[key]
    slot[0] += 1
    if wpis in kolumny_nieobecnosci:
        slot[1] += 1

with open(outfile, "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(["uczen_id", "Dziennik", "semestr", "wszystkich", "nieobecnych"])
    for (h, dziennik, semestr), (wszystkich, nieobecnych) in agg.items():
        w.writerow([h, dziennik, semestr, wszystkich, nieobecnych])

print(f"OK {rok}: {len(agg)} wierszy uczen-klasa-semestr z {i+1} wierszy zrodlowych, {len(hash_cache)} unikalnych uczniow -> {outfile}")
